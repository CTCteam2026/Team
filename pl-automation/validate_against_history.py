#!/usr/bin/env python3
"""
Back-test: feed each historical event's real inputs and real scope selection into
the builder, then compare against the P&L that actually went out the door.

Run AFTER build_pl_template.py.  Requires LibreOffice.
"""
import json, shutil, subprocess, sys
import openpyxl

SRC = "CTC_Conference_PL_Builder.xlsx"
RECALC = "/root/.claude/skills/synced/xlsx/scripts/recalc.py"
SCOPE_FIRST, SCOPE_LAST = 4, 76
INC_COL, LABEL_COL, HEAD_COL = 4, 2, 13

# INPUTS cells
K = dict(att="C10", exh="C11", spo="C12", spk="C13", ven="C14", sit="C15",
         evd="C18", setd="C19", trv="C20", mon="C22",
         onm="C27", onc="C28", prm="C29", prc="C30", mgrhrs="C42")

CASES = [
 dict(key="LCT", name="LCT / Marine Recreation Assn 2026",
      inputs={K["att"]: 200, K["exh"]: 35, K["spo"]: 0, K["spk"]: 25,
              K["ven"]: 6, K["sit"]: 1, K["evd"]: 3, K["setd"]: 0,
              K["trv"]: 0, K["mon"]: 4, K["onm"]: 2, K["onc"]: 2,
              K["prm"]: 1, K["prc"]: 1, K["mgrhrs"]: 14, "C44": 0},
      # LCT billed site visits inside its Venue Management line, and its
      # "Third Party Vendors" line covered the photographer.
      scope=["VENUE MANAGEMENT:",
             "EVENT BRANDING: Attendee Experience", "EVENT BRANDING: Signage",
             "EVENT BRANDING: Décor", "FOOD & BEVERAGE:",
             "REGISTRATION: Onsite Registration", "THIRD PARTY VENDORS: Vendors",
             "STAFF MANAGEMENT:", "SPEAKER MANAGEMENT:",
             "PROGRAMMING: Audio Visual", "PROGRAMMING: Entertainment & Talent",
             "RECEPTION PLANNING:", "EXHIBIT MANAGEMENT:", "POST-EVENT:"],
      actual=dict(pre_hrs=263, pre_client=22315, on_hrs=132, on_client=11100,
                  subtotal=33415)),
 dict(key="AFCI", name="AFCI Studio Summit 2027",
      inputs={K["att"]: 500, K["exh"]: 65, K["spo"]: 19, K["spk"]: 0,
              K["ven"]: 5, K["sit"]: 0, K["evd"]: 3, K["setd"]: 1,
              K["trv"]: 0, K["mon"]: 9, K["onm"]: 4, K["onc"]: 2,
              K["prm"]: 1, K["prc"]: 1, K["mgrhrs"]: 12},
      scope=["VENUE MANAGEMENT:",
             "EVENT BRANDING: Attendee Experience", "EVENT BRANDING: Signage",
             "EVENT BRANDING: Graphic Design", "REGISTRATION:",
             "THIRD PARTY VENDORS: Vendors", "FINANCIAL MANAGEMENT:",
             "SPONSORSHIPS:", "PROGRAMMING: Audio Visual",
             "PROGRAMMING: Stage Production",
             "PROGRAMMING: Recording & Broadcast — Onsite Support",
             "EXHIBIT MANAGEMENT:", "POST-EVENT:"],
      actual=dict(pre_hrs=617, pre_client=51845, on_hrs=272, on_client=27800,
                  subtotal=79645)),
 dict(key="WTUI", name="WTUI 2027",
      inputs={K["att"]: 1500, K["exh"]: 200, K["spo"]: 50, K["spk"]: 0,
              K["ven"]: 10, K["sit"]: 2, K["evd"]: 4, K["setd"]: 1,
              K["trv"]: 0, K["mon"]: 9, K["onm"]: 4, K["onc"]: 3,
              K["prm"]: 1, K["prc"]: 1, K["mgrhrs"]: 10},
      scope=["VENUE SOURCING:",
             "VENUE MANAGEMENT:", "EVENT BRANDING:", "FOOD & BEVERAGE:",
             "REGISTRATION:", "HOUSING LOGISTICS:",
             "THIRD PARTY VENDORS: Vendors", "STAFF MANAGEMENT:",
             "SPONSORSHIPS:", "PROGRAMMING: Audio Visual",
             "PROGRAMMING: Stage Production", "EVENT MARKETING:",
             "EXHIBIT MANAGEMENT:", "POST-EVENT:", "TOURNAMENTS:"],
      actual=dict(pre_hrs=1182, pre_client=107870, on_hrs=350, on_client=30725,
                  subtotal=138595)),
]


def run_case(case, verbose=False):
    path = f"backtest_{case['key']}.xlsx"
    shutil.copy(SRC, path)
    wb = openpyxl.load_workbook(path)
    iw, sc = wb["INPUTS"], wb["SCOPE"]
    iw["C5"] = case["name"]
    for cell, val in case["inputs"].items():
        iw[cell] = val
    for r in range(SCOPE_FIRST, SCOPE_LAST + 1):
        head = sc.cell(row=r, column=HEAD_COL).value
        if not head:                       # banded heading row, nothing to set
            continue
        cur = sc.cell(row=r, column=INC_COL)
        if "always included" in (sc.cell(row=r, column=5).value or ""):
            continue                       # always-on line
        label = f"{head}: {(sc.cell(row=r, column=LABEL_COL).value or '').strip()}"
        cur.value = "Yes" if any(label.startswith(p) for p in case["scope"]) else "No"
    wb.save(path)

    res = json.loads(subprocess.run([sys.executable, RECALC, path, "220"],
                                    capture_output=True, text=True).stdout)
    if res.get("status") != "success":
        raise SystemExit(f"{case['key']}: {res}")

    v = openpyxl.load_workbook(path, data_only=True)
    p = v["P&L"]
    if verbose:
        s = v["SCOPE"]
        for r in range(4, 24):
            if (s.cell(row=r, column=16).value or 0) > 0:
                print(f"    {s.cell(row=r, column=15).value:<40}"
                      f"{s.cell(row=r, column=16).value:>6.0f} hrs")
    return dict(pre_hrs=p["D43"].value, pre_client=p["E43"].value,
                on_hrs=p["D64"].value, on_client=p["E64"].value,
                subtotal=p["C9"].value, grand=p["C12"].value,
                cost=p["D12"].value, profit=p["E12"].value,
                margin=p["F12"].value, errors=res["total_errors"])


def d(m, a):
    return f"{(m - a) / a:+.0%}" if a else "n/a"


verbose = "-v" in sys.argv
print(f"{'EVENT':<30}{'METRIC':<18}{'MODEL':>10}{'ACTUAL':>10}{'DELTA':>8}")
print("-" * 76)
for case in CASES:
    m = run_case(case, verbose)
    a = case["actual"]
    for k, lbl in (("pre_hrs", "pre-planning hrs"), ("pre_client", "pre-planning $"),
                   ("on_hrs", "onsite hrs"), ("on_client", "onsite $"),
                   ("subtotal", "client subtotal $")):
        print(f"{case['name'][:29]:<30}{lbl:<18}{m[k]:>10,.0f}{a[k]:>10,.0f}{d(m[k], a[k]):>8}")
    print(f"{'':<30}{'grand total':<18}{m['grand']:>10,.0f}{'':>10}{'':>8}"
          f"   cost {m['cost']:,.0f} · profit {m['profit']:,.0f} · "
          f"margin {m['margin']:.1%} · errors {m['errors']}")
    print("-" * 76)
